"""Build anonymous review documents from the frozen content sources.

Content is authoritative. Tables stay editable and figures carry SVG plus PNG.
No legacy file is imported or modified.
"""
from pathlib import Path
import re
import json
import hashlib
import zipfile
from lxml import etree as ET
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def plain(s):
    s = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', s)
    return s.replace('**', '').replace('*', '').replace('`','')

def inline(p, s):
    pat = re.compile(r'\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|\*([^*]+)\*|`([^`]+)`')
    pos = 0
    for match in pat.finditer(s):
        p.add_run(s[pos:match.start()])
        if match[1] is not None:
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), p.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            text = OxmlElement('w:t'); text.text = match[1]
            run.append(text); link.append(run); p._p.append(link)
        elif match[5] is not None:
            run=p.add_run(match[5]);run.font.name='Consolas';run.font.size=Pt(9)
        else:
            run = p.add_run(match[3] or match[4])
            run.bold = bool(match[3]); run.italic = bool(match[4])
        pos = match.end()
    p.add_run(s[pos:])

def border(parent, group, side, size='5', value='single'):
    b = parent.find(qn('w:' + group))
    if b is None:
        b = OxmlElement('w:' + group); parent.append(b)
    el = OxmlElement('w:' + side)
    for key, val in [('val',value),('sz',size),('color','404040')]:
        el.set(qn('w:' + key), val)
    b.append(el)


def column_widths(rows, total=6.75):
    """Proportional fixed widths so long text columns stop collapsing."""
    n = len(rows[0])
    weight = [max(len(r[i]) for r in rows) ** 0.75 for i in range(n)]
    scale = total / sum(weight)
    widths = [max(0.75, w * scale) for w in weight]
    scale = total / sum(widths)
    return [round(w * scale, 2) for w in widths]

def add_table(doc, rows):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    props = table._tbl.tblPr
    for side in ['top', 'bottom']:
        border(props, 'tblBorders', side, '7')
    for side in ['left','right','insideH','insideV']:
        border(props, 'tblBorders', side, '0', 'nil')
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        trp = table.rows[-1]._tr.get_or_add_trPr()
        trp.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trp.append(OxmlElement('w:tblHeader'))
        for cell, value in zip(cells,row):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; inline(p,value)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.line_spacing = 1.05
            if i == 0 or (len(rows) <= 3 and i < len(rows) - 1):
                p.paragraph_format.keep_with_next = True
            for run in p.runs:
                run.font.name = 'Arial'; run.font.size = Pt(9.5)
                if i == 0: run.bold = True
            if i == 0: border(cell._tc.get_or_add_tcPr(), 'tcBorders', 'bottom')
    widths = column_widths(rows)
    table.autofit = False
    table.allow_autofit = False
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed')
    props.append(layout)
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def native_svgs(path, figures):
    if not figures: return []
    ns = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
          'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
          'pr':'http://schemas.openxmlformats.org/package/2006/relationships',
          'ct':'http://schemas.openxmlformats.org/package/2006/content-types',
          'asvg':'http://schemas.microsoft.com/office/drawing/2016/SVG/main'}
    with zipfile.ZipFile(path) as archive:
        contents = {n:archive.read(n) for n in archive.namelist()}
    document = ET.fromstring(contents['word/document.xml'])
    rels = ET.fromstring(contents['word/_rels/document.xml.rels'])
    types = ET.fromstring(contents['[Content_Types].xml'])
    blips = document.xpath('//a:blip', namespaces=ns)
    assert len(blips) == len(figures)
    audit = []
    for i,(blip,png) in enumerate(zip(blips,figures),1):
        svg = png.with_suffix('.svg'); assert svg.is_file()
        ET.fromstring(svg.read_bytes())
        part = f'media/evidence_audit_{i}.svg'; rid = f'rIdAuditSVG{i}'
        contents['word/' + part] = svg.read_bytes()
        rel = ET.SubElement(rels, '{'+ns['pr']+'}Relationship')
        rel.set('Id',rid); rel.set('Type',RT.IMAGE); rel.set('Target',part)
        extlist = ET.SubElement(blip, '{'+ns['a']+'}extLst')
        ext = ET.SubElement(extlist,'{'+ns['a']+'}ext')
        ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        svgblip = ET.SubElement(ext,'{'+ns['asvg']+'}svgBlip',nsmap={'asvg':ns['asvg']})
        svgblip.set('{'+ns['r']+'}embed',rid)
        audit.append({'path':str(svg),'sha256':sha(svg),'part':'word/'+part})
    if not types.xpath('./ct:Default[@Extension="svg"]',namespaces=ns):
        d = ET.SubElement(types,'{'+ns['ct']+'}Default')
        d.set('Extension','svg');d.set('ContentType','image/svg+xml')
    for name,el in [('word/document.xml',document),('word/_rels/document.xml.rels',rels),('[Content_Types].xml',types)]:
        contents[name] = ET.tostring(el,encoding='UTF-8',xml_declaration=True,standalone=True)
    temp = path.with_suffix('.patch.tmp')
    with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as archive:
        for name,payload in contents.items(): archive.writestr(name,payload)
    temp.replace(path)
    return audit

def build(source, output):
    content = source.read_text(encoding='utf-8')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5);sec.page_height = Inches(11)
    sec.top_margin = Inches(.8);sec.bottom_margin = Inches(.8)
    sec.left_margin = Inches(.85);sec.right_margin = Inches(.85)
    for name,size in [('Normal',11),('Title',17),('Heading 1',13),('Heading 2',11.5)]:
        st = doc.styles[name];st.font.name = 'Times New Roman';st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0,0,0)
        if name != 'Normal': st.font.bold = True
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = 1.12;normal.space_after = Pt(7)
    paragraphs=[]; tables=[]; figures=[]
    lines = content.splitlines();i=0
    while i < len(lines):
        line = lines[i].strip();i+=1
        if not line: continue
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip());i+=1
            rows=[[v.strip() for v in row.strip('|').split('|')] for row in block]
            rows=[row for row in rows if not all(re.fullmatch(r':?-+:?',v) for v in row)]
            assert all(len(row)==len(rows[0]) for row in rows)
            add_table(doc,rows);tables.append([[plain(v) for v in row] for row in rows]);continue
        m=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if m:
            path=(source.parent/m[2]).resolve();assert path.is_file()
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(path),width=Inches(6.8))
            figures.append(path);continue
        heading=re.match(r'^(#{1,3}) (.+)$',line)
        if heading:
            style={1:'Title',2:'Heading 1',3:'Heading 2'}[len(heading[1])]
            p=doc.add_paragraph(plain(heading[2]),style=style);paragraphs.append(plain(heading[2]));continue
        quoted=line.startswith('> ')
        if quoted:line=line[2:]
        p=doc.add_paragraph();inline(p,line);paragraphs.append(plain(line))
        if quoted:
            p.paragraph_format.left_indent=Inches(.18)
            p.paragraph_format.right_indent=Inches(.12)
        next_content = next((item.strip() for item in lines[i:] if item.strip()), '')
        if re.match(r'Table [A-Z]?\d+\b',line) and next_content.startswith('|'):
            p.paragraph_format.keep_with_next=True
            for run in p.runs:run.bold=True
        if re.match(r'Figure \d+\b',line):
            p.paragraph_format.keep_together=True
            p.paragraph_format.keep_with_next=False
    doc.core_properties.title=plain(lines[0][2:])
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.comments=''
    for el in [doc.element,doc.styles.element]:
        for b in el.xpath('.//w:pBdr'):b.getparent().remove(b)
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    svgs=native_svgs(output,figures)
    check=Document(output)
    actual=[''.join(p._p.xpath('.//w:t/text()')) for p in check.paragraphs]
    actual=[p for p in actual if p]
    assert actual==paragraphs,('Paragraph mismatch',source)
    assert [[[c.text for c in r.cells] for r in t.rows] for t in check.tables]==tables
    assert len(check.inline_shapes)==len(figures)
    return {'source':str(source),'source_sha256':sha(source),'output':str(output),'sha256':sha(output),
            'paragraphs_exact':True,'table_cells_exact':True,'tables':len(tables),
            'cells':sum(sum(len(row) for row in tb) for tb in tables),'figures':len(figures),
            'native_svg':svgs,'render_required':True}

def main():
    targets=[('manuscript/manuscript.md','deliverables/Tariff_Evidence_Audit_20261023.docx'),
             ('manuscript/supplement.md','deliverables/Tariff_Evidence_Audit_Supplement_20261023.docx')]
    reports=[build(ROOT/s,ROOT/o) for s,o in targets]
    (ROOT/'verification/docx_content_audit.json').write_text(json.dumps(reports,indent=2),encoding='utf-8')
    print(json.dumps(reports,indent=2))

if __name__=='__main__':main()
